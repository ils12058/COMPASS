from __future__ import annotations

from datetime import date
from io import BytesIO

import pytest
from django.core.management import call_command
from django.utils import timezone
from openpyxl import load_workbook

from compass.accounts.models import Role, StudentLifecycleStatus, User
from compass.graduate_tracer.models import (
    GraduateTracerEducation,
    GraduateTracerResponse,
    GraduateTracerStatus,
    GTSBusinessLine,
    GTSEmploymentState,
    GTSPresentEmploymentStatus,
)
from compass.reports import graduate_tracer_xlsx as report_xlsx
from compass.reports.graduate_tracer_xlsx import (
    EMPTY_REPORT_MESSAGE,
    render_graduate_tracer_xlsx,
)
from tests.graduate_tracer_test_helpers import encrypted_row


def sync_policy() -> None:
    call_command("sync_identity_policy", verbosity=0)


def make_student(email: str) -> User:
    user = User.objects.create_user(
        email=email,
        password="a-test-password",
        role=Role.objects.get(code="STUDENT"),
        first_name="Graduate",
        last_name="Workbook",
    )
    user.student_lifecycle_status = StudentLifecycleStatus.GRADUATED
    user.save(update_fields=["student_lifecycle_status", "updated_at"])
    return user


def make_response(email: str, **values) -> GraduateTracerResponse:
    defaults = {
        "sex": "MALE",
        "civil_status": "SINGLE",
        "region_of_origin": "REGION_5",
        "residence_location": "MUNICIPALITY",
        "current_employment_state": GTSEmploymentState.NOT_EMPLOYED,
        "unemployment_reasons": ["NO_JOB_OPPORTUNITY"],
    }
    defaults.update(values)
    return encrypted_row(
        GraduateTracerResponse,
        student=make_student(email),
        status=GraduateTracerStatus.SUBMITTED,
        submitted_at=timezone.now(),
        **defaults,
    )


def workbook_values(workbook) -> list[str]:
    values: list[str] = []
    for worksheet in workbook.worksheets:
        for row in worksheet.iter_rows():
            for cell in row:
                if cell.value is not None:
                    values.append(str(cell.value))
    return values


@pytest.mark.django_db
def test_graduate_tracer_xlsx_is_aggregate_only_safe_and_compact():
    sync_policy()
    sentinel = "SENTINEL-RAW-GTS-XLSX"
    item = make_response(
        "gts-xlsx-private@example.edu",
        name_snapshot=sentinel,
        permanent_address_snapshot=sentinel,
        present_occupation=sentinel,
        unemployment_other_reason=sentinel,
        curriculum_improvement_suggestions=sentinel,
    )
    encrypted_row(
        GraduateTracerEducation,
        response=item,
        position=1,
        degree_and_specialization=sentinel,
        college_or_university=sentinel,
        year_graduated=2026,
        honors_or_awards=sentinel,
    )
    make_response(
        "gts-xlsx-employed@example.edu",
        current_employment_state=GTSEmploymentState.EMPLOYED,
        unemployment_reasons=[],
        present_employment_status=GTSPresentEmploymentStatus.REGULAR_PERMANENT,
        employer_business_line=GTSBusinessLine.EDUCATION,
        place_of_work="LOCAL",
        first_job_after_college=False,
        first_job_duration="ONE_TO_SIX_MONTHS",
        first_job_source="ADVERTISEMENT",
        time_to_first_job="LESS_THAN_MONTH",
        first_job_level="PROFESSIONAL_TECHNICAL_SUPERVISORY",
        current_job_level="PROFESSIONAL_TECHNICAL_SUPERVISORY",
        initial_gross_monthly_earning="FROM_15000_TO_LT_20000",
        curriculum_relevant_to_first_job=False,
    )

    result = render_graduate_tracer_xlsx()
    assert result.xlsx_bytes.startswith(b"PK")
    assert result.filename == "graduate-tracer-schema-v1.xlsx"

    workbook = load_workbook(BytesIO(result.xlsx_bytes), data_only=False)
    assert workbook.sheetnames == [
        "Summary",
        "Respondent Profile",
        "Employment",
        "First Job",
        "Reasons & Skills",
    ]
    assert all(worksheet.sheet_state == "visible" for worksheet in workbook.worksheets)
    summary = workbook["Summary"]
    assert summary["A3"].value == "At a Glance"
    assert summary["A4"].value == "Submitted Response Count"
    assert summary["B4"].value == 2
    assert summary["B4"].data_type == "n"
    assert summary.freeze_panes == "A4"

    values = workbook_values(workbook)
    serialized = "\n".join(values)
    assert sentinel not in serialized
    assert "gts-xlsx-private@example.edu" not in serialized
    assert "Raw" not in workbook.sheetnames

    for worksheet in workbook.worksheets:
        assert worksheet.page_setup.orientation == "landscape"
        assert str(worksheet.page_setup.paperSize) == str(worksheet.PAPERSIZE_A4)
        assert worksheet.page_setup.fitToWidth == 1
        assert worksheet.print_area
        for row in worksheet.iter_rows():
            for cell in row:
                assert cell.data_type != "f"
                assert cell.hyperlink is None

    employment = workbook["Employment"]
    employed_row = None
    for row in employment.iter_rows():
        if row[0].value == GTSEmploymentState.EMPLOYED.label:
            employed_row = row
            break
    assert employed_row is not None
    assert employed_row[1].value == 1
    assert float(employed_row[2].value) == 50.0
    assert employed_row[2].number_format == "0.00"


@pytest.mark.django_db
def test_long_aggregate_label_wraps_without_clipping(monkeypatch):
    sync_policy()
    original = report_xlsx.build_graduate_tracer_report
    long_label = "Synthetic long category label " * 5

    def report_with_long_label(**kwargs):
        report = original(**kwargs)
        report["sections"]["sex"]["rows"][0]["label"] = long_label
        return report

    monkeypatch.setattr(report_xlsx, "build_graduate_tracer_report", report_with_long_label)
    result = report_xlsx.render_graduate_tracer_xlsx()
    workbook = load_workbook(BytesIO(result.xlsx_bytes), data_only=False)
    worksheet = workbook["Respondent Profile"]
    row = next(row for row in worksheet.iter_rows() if row[0].value == long_label)
    assert row[0].alignment.wrap_text is True
    assert worksheet.row_dimensions[row[0].row].height > 18


@pytest.mark.django_db
def test_zero_result_xlsx_still_has_context_methodology_and_zero_sections():
    result = render_graduate_tracer_xlsx()
    workbook = load_workbook(BytesIO(result.xlsx_bytes), data_only=False)

    values = workbook_values(workbook)
    assert EMPTY_REPORT_MESSAGE in values
    assert "Submitted Response Count" in values
    assert "0" in values
    assert any("No response rate is calculated" in value for value in values)
    assert "Small-population privacy notice" not in values
    assert "N/A" in values
    assert workbook.sheetnames == [
        "Summary",
        "Respondent Profile",
        "Employment",
        "First Job",
        "Reasons & Skills",
    ]


@pytest.mark.django_db
def test_filtered_xlsx_filename_is_deterministic_and_contains_only_safe_dates():
    result = render_graduate_tracer_xlsx(
        submitted_from=date(2026, 1, 1),
        submitted_to=date(2026, 12, 31),
    )
    assert result.filename == "graduate-tracer-schema-v1-submitted-2026-01-01-to-2026-12-31.xlsx"
    assert result.release_context == {
        "instrument_schema_version": 1,
        "submitted_from": "2026-01-01",
        "submitted_to": "2026-12-31",
    }


@pytest.mark.django_db
def test_small_population_xlsx_includes_privacy_warning_and_keeps_exact_values():
    sync_policy()
    make_response("gts-xlsx-small@example.edu")

    result = render_graduate_tracer_xlsx()
    workbook = load_workbook(BytesIO(result.xlsx_bytes), data_only=False)
    values = workbook_values(workbook)

    assert "Small-population privacy notice" in values
    assert any("Use and share these results carefully" in value for value in values)
    profile = workbook["Respondent Profile"]
    male_row = next(row for row in profile.iter_rows() if row[0].value == "Male")
    assert male_row[1].value == 1
    assert float(male_row[2].value) == 100.0
