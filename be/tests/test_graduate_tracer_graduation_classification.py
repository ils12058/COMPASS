"""Tests for GTS-C1 Graduate Tracer graduation classification and academic filtering."""

from __future__ import annotations

import json
from datetime import date
from decimal import Decimal
from uuid import uuid4

import pytest
from django.utils import timezone
from openpyxl import load_workbook
from io import BytesIO

from compass.accounts.models import Designation, Role, StudentLifecycleStatus, User, UserDesignation
from compass.common.institutional_time import institution_today
from compass.graduate_tracer.disposition import anonymize_response, verify_anonymized
from compass.graduate_tracer.models import (
    GraduateTracerResponse,
    GraduateTracerStatus,
    GTSEmploymentState,
)
from compass.graduate_tracer.services import (
    ensure_my_response,
    replace_my_draft,
    submit_my_response,
)
from compass.organization.models import Campus, College, Program
from compass.reports.graduate_tracer import (
    GraduateTracerDisclosureRestricted,
    InvalidGraduateTracerReportFilter,
    build_graduate_tracer_report,
)
from compass.reports.graduate_tracer_xlsx import render_graduate_tracer_xlsx
from tests.test_graduate_tracer import (
    auth_client,
    csrf,
    default_academic_classification,
    make_head,
    make_user,
    post_empty,
    put_json,
    sync_policy,
    valid_employed_payload,
    valid_unemployed_payload,
)


def create_org_hierarchy():
    campus = Campus.objects.create(code="MAIN-UCN", name="Main Campus")
    college_a = College.objects.create(campus=campus, code="CCMS", name="College of Computing")
    college_b = College.objects.create(campus=campus, code="CBA", name="College of Business")
    prog_a1 = Program.objects.create(college=college_a, code="BSIS", name="BS Information Systems")
    prog_a2 = Program.objects.create(college=college_a, code="BSCS", name="BS Computer Science")
    prog_b1 = Program.objects.create(college=college_b, code="BSBA", name="BS Business Administration")
    return college_a, college_b, prog_a1, prog_a2, prog_b1


@pytest.mark.django_db
def test_classification_options_endpoint_permissions_and_response():
    sync_policy()
    college_a, college_b, prog_a1, prog_a2, prog_b1 = create_org_hierarchy()
    inactive_college = College.objects.create(
        campus=college_a.campus, code="INACT", name="Inactive College", is_active=False
    )
    Program.objects.create(
        college=inactive_college, code="INACT-P", name="Inactive Program", is_active=False
    )

    student = make_user("student-options@example.edu")
    head = make_head("head-options@example.edu")
    it_admin = make_user("admin-options@example.edu", role="IT_ADMIN", lifecycle=None)

    # Student with manage_self / view_self capability
    student_res = auth_client(student).get("/api/v1/graduate-tracer/classification-options")
    assert student_res.status_code == 200
    data = student_res.json()
    assert "colleges" in data and "programs" in data
    college_codes = [c["code"] for c in data["colleges"]]
    program_codes = [p["code"] for p in data["programs"]]
    assert "CCMS" in college_codes and "CBA" in college_codes
    assert "INACT" not in college_codes
    assert "BSIS" in program_codes and "BSBA" in program_codes
    assert "INACT-P" not in program_codes

    # Program option links to college_id
    bsis_option = next(p for p in data["programs"] if p["code"] == "BSIS")
    assert bsis_option["college_id"] == str(college_a.pk)

    # Head guidance with graduate_tracer.view / reports.view
    assert auth_client(head).get("/api/v1/graduate-tracer/classification-options").status_code == 200

    # IT Admin without GTS capabilities
    assert auth_client(it_admin).get("/api/v1/graduate-tracer/classification-options").status_code == 403


@pytest.mark.django_db
def test_draft_allows_missing_classification_submission_enforces_it():
    sync_policy()
    college_a, college_b, prog_a1, prog_a2, prog_b1 = create_org_hierarchy()
    student = make_user("student-draft-class@example.edu")
    client = auth_client(student)
    assert post_empty(client, "/api/v1/graduate-tracer/me").status_code == 200

    payload = valid_employed_payload()
    # Draft can be saved with no academic classification
    payload["graduated_college_id"] = None
    payload["graduated_program_id"] = None
    payload["graduation_year"] = None
    draft_res = put_json(client, "/api/v1/graduate-tracer/me", payload)
    assert draft_res.status_code == 200
    assert draft_res.json()["graduated_college"] is None
    assert draft_res.json()["graduated_program"] is None
    assert draft_res.json()["graduation_year"] is None

    # Submitting with missing classification fails (422)
    assert post_empty(client, "/api/v1/graduate-tracer/me/submit").status_code == 422

    # Partial classification: college only
    payload["graduated_college_id"] = str(college_a.pk)
    put_json(client, "/api/v1/graduate-tracer/me", payload)
    assert post_empty(client, "/api/v1/graduate-tracer/me/submit").status_code == 422

    # Partial classification: college + program, missing year
    payload["graduated_program_id"] = str(prog_a1.pk)
    put_json(client, "/api/v1/graduate-tracer/me", payload)
    assert post_empty(client, "/api/v1/graduate-tracer/me/submit").status_code == 422

    # Incompatible program-college combination
    payload["graduated_college_id"] = str(college_b.pk)
    payload["graduated_program_id"] = str(prog_a1.pk)  # prog_a1 belongs to college_a
    payload["graduation_year"] = 2024
    put_json(client, "/api/v1/graduate-tracer/me", payload)
    assert post_empty(client, "/api/v1/graduate-tracer/me/submit").status_code == 422

    # Future graduation year
    future_year = institution_today().year + 1
    payload["graduated_college_id"] = str(college_a.pk)
    payload["graduated_program_id"] = str(prog_a1.pk)
    payload["graduation_year"] = future_year
    put_json(client, "/api/v1/graduate-tracer/me", payload)
    assert post_empty(client, "/api/v1/graduate-tracer/me/submit").status_code == 422

    # Valid classification: college_a, prog_a1, 2024
    payload["graduation_year"] = 2024
    put_res = put_json(client, "/api/v1/graduate-tracer/me", payload)
    assert put_res.status_code == 200
    assert put_res.json()["graduated_college"]["code"] == "CCMS"
    assert put_res.json()["graduated_program"]["code"] == "BSIS"
    assert put_res.json()["graduation_year"] == 2024

    submit_res = post_empty(client, "/api/v1/graduate-tracer/me/submit")
    assert submit_res.status_code == 200
    assert submit_res.json()["status"] == "SUBMITTED"
    assert submit_res.json()["graduated_college"]["code"] == "CCMS"
    assert submit_res.json()["graduated_program"]["code"] == "BSIS"
    assert submit_res.json()["graduation_year"] == 2024


@pytest.mark.django_db
def test_roster_filtering_and_contradictory_combination_rejection():
    sync_policy()
    college_a, college_b, prog_a1, prog_a2, prog_b1 = create_org_hierarchy()
    head = make_head("head-roster-test@example.edu")
    client = auth_client(head)

    # Student 1: College A, Prog A1, 2023
    s1 = make_user("s1-class@example.edu")
    c1 = auth_client(s1)
    post_empty(c1, "/api/v1/graduate-tracer/me")
    p1 = valid_employed_payload()
    p1["graduated_college_id"] = str(college_a.pk)
    p1["graduated_program_id"] = str(prog_a1.pk)
    p1["graduation_year"] = 2023
    put_json(c1, "/api/v1/graduate-tracer/me", p1)
    post_empty(c1, "/api/v1/graduate-tracer/me/submit")

    # Student 2: College A, Prog A2, 2024
    s2 = make_user("s2-class@example.edu")
    c2 = auth_client(s2)
    post_empty(c2, "/api/v1/graduate-tracer/me")
    p2 = valid_unemployed_payload()
    p2["graduated_college_id"] = str(college_a.pk)
    p2["graduated_program_id"] = str(prog_a2.pk)
    p2["graduation_year"] = 2024
    put_json(c2, "/api/v1/graduate-tracer/me", p2)
    post_empty(c2, "/api/v1/graduate-tracer/me/submit")

    # Student 3: College B, Prog B1, 2024
    s3 = make_user("s3-class@example.edu")
    c3 = auth_client(s3)
    post_empty(c3, "/api/v1/graduate-tracer/me")
    p3 = valid_employed_payload()
    p3["graduated_college_id"] = str(college_b.pk)
    p3["graduated_program_id"] = str(prog_b1.pk)
    p3["graduation_year"] = 2024
    put_json(c3, "/api/v1/graduate-tracer/me", p3)
    post_empty(c3, "/api/v1/graduate-tracer/me/submit")

    # Unfiltered: all 3
    all_res = client.get("/api/v1/graduate-tracer/responses")
    assert all_res.status_code == 200
    assert len(all_res.json()["items"]) == 3

    # Filter by College A: s1 and s2
    col_a_res = client.get(f"/api/v1/graduate-tracer/responses?college_id={college_a.pk}")
    assert col_a_res.status_code == 200
    assert len(col_a_res.json()["items"]) == 2

    # Filter by Program A1: s1 only
    prog_a1_res = client.get(f"/api/v1/graduate-tracer/responses?program_id={prog_a1.pk}")
    assert prog_a1_res.status_code == 200
    assert len(prog_a1_res.json()["items"]) == 1
    assert prog_a1_res.json()["items"][0]["student_id"] == str(s1.pk)

    # Filter by Graduation Year 2024: s2 and s3
    year_2024_res = client.get("/api/v1/graduate-tracer/responses?graduation_year=2024")
    assert year_2024_res.status_code == 200
    assert len(year_2024_res.json()["items"]) == 2

    # Filter by College A + Year 2024: s2 only
    col_year_res = client.get(
        f"/api/v1/graduate-tracer/responses?college_id={college_a.pk}&graduation_year=2024"
    )
    assert col_year_res.status_code == 200
    assert len(col_year_res.json()["items"]) == 1
    assert col_year_res.json()["items"][0]["student_id"] == str(s2.pk)

    # Contradictory College + Program: College B + Prog A1 -> 422
    contra_res = client.get(
        f"/api/v1/graduate-tracer/responses?college_id={college_b.pk}&program_id={prog_a1.pk}"
    )
    assert contra_res.status_code == 422


@pytest.mark.django_db
def test_aggregate_report_filtering_and_disclosure_restrictions():
    sync_policy()
    college_a, college_b, prog_a1, prog_a2, prog_b1 = create_org_hierarchy()
    head = make_head("head-report-class@example.edu")
    client = auth_client(head)

    # 1. Zero matching responses under filter -> standard empty report (200), not restricted
    empty_res = client.get(
        f"/api/v1/reports/graduate-tracer?college_id={college_a.pk}&graduation_year=2020"
    )
    assert empty_res.status_code == 200
    assert empty_res.json()["report_context"]["submitted_response_count"] == 0

    # 2. Create 2 responses in College A / Prog A1 / 2024
    for i in range(2):
        u = make_user(f"small-cohort-{i}@example.edu")
        uc = auth_client(u)
        post_empty(uc, "/api/v1/graduate-tracer/me")
        p = valid_employed_payload()
        p["graduated_college_id"] = str(college_a.pk)
        p["graduated_program_id"] = str(prog_a1.pk)
        p["graduation_year"] = 2024
        put_json(uc, "/api/v1/graduate-tracer/me", p)
        post_empty(uc, "/api/v1/graduate-tracer/me/submit")

    # Filtered academic query matches 2 (0 < count < 5) -> 403 small_population_disclosure_restricted
    restricted_res = client.get(
        f"/api/v1/reports/graduate-tracer?college_id={college_a.pk}&graduation_year=2024"
    )
    assert restricted_res.status_code == 403
    assert restricted_res.json()["error"]["code"] == "small_population_disclosure_restricted"

    # XLSX export also fails-closed with 403
    xlsx_res = client.get(
        f"/api/v1/reports/graduate-tracer/xlsx?college_id={college_a.pk}&graduation_year=2024"
    )
    assert xlsx_res.status_code == 403
    assert xlsx_res.json()["error"]["code"] == "small_population_disclosure_restricted"

    # Unfiltered aggregate with same 2 responses returns 200 with advisory warning
    unfiltered_res = client.get("/api/v1/reports/graduate-tracer")
    assert unfiltered_res.status_code == 200
    assert unfiltered_res.json()["report_context"]["submitted_response_count"] == 2
    assert [w["code"] for w in unfiltered_res.json()["disclosure_warnings"]] == ["SMALL_POPULATION"]

    # 3. Add 3 more responses to reach threshold of 5
    for i in range(2, 5):
        u = make_user(f"reach-threshold-{i}@example.edu")
        uc = auth_client(u)
        post_empty(uc, "/api/v1/graduate-tracer/me")
        p = valid_employed_payload()
        p["graduated_college_id"] = str(college_a.pk)
        p["graduated_program_id"] = str(prog_a1.pk)
        p["graduation_year"] = 2024
        put_json(uc, "/api/v1/graduate-tracer/me", p)
        post_empty(uc, "/api/v1/graduate-tracer/me/submit")

    # Filtered academic query now matches 5 (>= 5) -> succeeds 200
    allowed_res = client.get(
        f"/api/v1/reports/graduate-tracer?college_id={college_a.pk}&graduation_year=2024"
    )
    assert allowed_res.status_code == 200
    body = allowed_res.json()
    assert body["report_context"]["submitted_response_count"] == 5
    assert body["report_context"]["college"]["code"] == "CCMS"
    assert body["report_context"]["graduation_year"] == 2024

    # XLSX export succeeds
    allowed_xlsx = client.get(
        f"/api/v1/reports/graduate-tracer/xlsx?college_id={college_a.pk}&graduation_year=2024"
    )
    assert allowed_xlsx.status_code == 200
    assert allowed_xlsx.headers["Content-Type"] == (
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    wb = load_workbook(BytesIO(allowed_xlsx.content))
    summary_sheet = wb["Summary"]
    summary_text = [cell.value for row in summary_sheet.iter_rows() for cell in row if cell.value]
    assert "College at Graduation" in summary_text
    assert "CCMS - College of Computing" in summary_text
    assert "Graduation Year" in summary_text
    assert "2024" in summary_text


@pytest.mark.django_db
def test_contradictory_report_filters_rejected_with_422():
    sync_policy()
    college_a, college_b, prog_a1, prog_a2, prog_b1 = create_org_hierarchy()
    head = make_head("head-contra-test@example.edu")
    client = auth_client(head)

    res = client.get(
        f"/api/v1/reports/graduate-tracer?college_id={college_b.pk}&program_id={prog_a1.pk}"
    )
    assert res.status_code == 422
    assert res.json()["error"]["code"] == "invalid_report_filter"

    xlsx_res = client.get(
        f"/api/v1/reports/graduate-tracer/xlsx?college_id={college_b.pk}&program_id={prog_a1.pk}"
    )
    assert xlsx_res.status_code == 422


@pytest.mark.django_db
def test_anonymization_retention_invariant_drops_quasi_identifiers():
    sync_policy()
    college_a, college_b, prog_a1, prog_a2, prog_b1 = create_org_hierarchy()
    student = make_user("anon-retention@example.edu")
    client = auth_client(student)
    post_empty(client, "/api/v1/graduate-tracer/me")
    payload = valid_employed_payload()
    payload["graduated_college_id"] = str(college_a.pk)
    payload["graduated_program_id"] = str(prog_a1.pk)
    payload["graduation_year"] = 2024
    put_json(client, "/api/v1/graduate-tracer/me", payload)
    post_empty(client, "/api/v1/graduate-tracer/me/submit")

    source = GraduateTracerResponse.objects.get(student=student)
    assert source.graduated_college_id == college_a.pk
    assert source.graduated_program_id == prog_a1.pk
    assert source.graduation_year == 2024

    # Perform operational anonymization
    anonymize_response(source.pk)

    # Anonymized response must have quasi-identifiers dropped (None)
    retained = GraduateTracerResponse.objects.get(anonymized_at__isnull=False)
    assert retained.student_id is None
    assert retained.graduated_college_id is None
    assert retained.graduated_program_id is None
    assert retained.graduation_year is None
    assert verify_anonymized(retained)
