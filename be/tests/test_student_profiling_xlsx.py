from __future__ import annotations

import copy
from datetime import date
from decimal import Decimal
from io import BytesIO
from uuid import uuid4

import pytest
from django.test import Client
from openpyxl import load_workbook

from compass.accounts.models import Designation, UserDesignation
from compass.inventory.models import CivilStatusCategory, Sex
from compass.organization.models import (
    AcademicYear,
    Campus,
    College,
    CounselorResponsibility,
    Program,
)
from compass.reports import api as reports_api
from compass.reports import xlsx as report_xlsx
from compass.reports.pdf import (
    build_student_profiling_print_context,
    student_profiling_pdf_filename,
)
from compass.reports.services import (
    GLOBAL_REPORT_ACCESS_SCOPE,
    InvalidReportFilter,
    ReportAccessScope,
    ReportConfigurationConflict,
    ReportNotFound,
    build_student_profiling_report,
)
from compass.reports.xlsx import (
    EMPTY_REPORT_MESSAGE,
    MAX_EXCEL_COLUMNS,
    SECTION_SPECS,
    XLSX_CONTENT_TYPE,
    StudentProfilingWorkbookUnavailable,
    StudentProfilingXlsxResult,
    _validate_excel_column_limit,
    build_student_profiling_workbook,
    render_student_profiling_xlsx,
    student_profiling_xlsx_filename,
)
from tests.test_student_profiling_pdf import (
    auth_client,
    make_head,
    make_user,
    sync_policy,
    synthetic_report,
)
from tests.test_student_profiling_reports import make_inventory, make_organization, make_revision
from tests.test_student_profiling_reports import make_user as make_profile_student

EXPECTED_SHEETS = [
    "Summary",
    "Sex",
    "Age",
    "Civil Status",
    "PWD Status",
    "Religion",
    "Mother Life",
    "Father Life",
    "Parent Family",
    "City Municipality",
    "Parent Income",
    "Mother Occupation",
    "Father Occupation",
    "Living Condition",
]


def fake_xlsx_result() -> StudentProfilingXlsxResult:
    return StudentProfilingXlsxResult(
        xlsx_bytes=b"synthetic-xlsx-payload",
        filename="student-profile-2026-2027.xlsx",
        release_context={
            "academic_year_id": str(uuid4()),
            "academic_year_label": "2026-2027",
            "campus_id": None,
            "campus_code": None,
            "college_id": None,
            "college_code": None,
            "program_id": None,
            "program_code": None,
            "year_level": None,
        },
    )


def workbook_from_report(monkeypatch, report):
    monkeypatch.setattr(
        report_xlsx,
        "build_student_profiling_report",
        lambda **kwargs: report,
    )
    result = render_student_profiling_xlsx(access_scope=GLOBAL_REPORT_ACCESS_SCOPE)
    return load_workbook(BytesIO(result.xlsx_bytes), data_only=False, keep_links=False)


def workbook_values(workbook) -> list[object]:
    return [
        cell.value
        for worksheet in workbook.worksheets
        for row in worksheet.iter_rows()
        for cell in row
        if cell.value is not None
    ]


def assert_no_formulas_or_hyperlinks(workbook) -> None:
    for worksheet in workbook.worksheets:
        for row in worksheet.iter_rows():
            for cell in row:
                assert cell.data_type != "f"
                assert cell.hyperlink is None


@pytest.mark.django_db
def test_xlsx_endpoint_reuses_reports_view_without_recent_mfa(monkeypatch):
    sync_policy()
    monkeypatch.setattr(
        reports_api,
        "render_student_profiling_xlsx",
        lambda **kwargs: fake_xlsx_result(),
    )

    head = make_head("head-xlsx@example.edu")
    counselor = make_user("ordinary-xlsx@example.edu", "COUNSELOR")
    staff = make_user("staff-xlsx@example.edu", "GUIDANCE_SERVICES_STAFF")
    student = make_user("student-xlsx@example.edu", "STUDENT")
    admin = make_user("admin-xlsx@example.edu", "IT_ADMIN")
    dpo = make_user("dpo-xlsx@example.edu", "INSTITUTIONAL_OFFICER")
    UserDesignation.objects.create(
        user=dpo,
        designation=Designation.objects.get(code="DPO"),
    )

    response = auth_client(head).get("/api/v1/reports/student-profile/xlsx")
    assert response.status_code == 200
    assert response["Content-Type"] == XLSX_CONTENT_TYPE
    assert response["Content-Disposition"] == (
        'attachment; filename="student-profile-2026-2027.xlsx"'
    )
    assert response.content == b"synthetic-xlsx-payload"

    for user in (counselor, staff, student, admin, dpo):
        assert auth_client(user).get("/api/v1/reports/student-profile/xlsx").status_code == 403
    assert Client().get("/api/v1/reports/student-profile/xlsx").status_code == 401


@pytest.mark.django_db
def test_xlsx_endpoint_forwards_same_report_filters(monkeypatch):
    sync_policy()
    head = make_head("head-xlsx-filters@example.edu")
    client = auth_client(head)
    academic_year_id = uuid4()
    campus_id = uuid4()
    college_id = uuid4()
    program_id = uuid4()
    captured: list[dict[str, object]] = []

    def fake_render(**kwargs):
        captured.append(kwargs)
        return fake_xlsx_result()

    monkeypatch.setattr(reports_api, "render_student_profiling_xlsx", fake_render)
    response = client.get(
        "/api/v1/reports/student-profile/xlsx",
        {
            "academic_year_id": academic_year_id,
            "campus_id": campus_id,
            "college_id": college_id,
            "program_id": program_id,
            "year_level": 4,
        },
    )
    assert response.status_code == 200
    assert captured == [
        {
            "academic_year_id": academic_year_id,
            "campus_id": campus_id,
            "college_id": college_id,
            "program_id": program_id,
            "year_level": 4,
            "access_scope": GLOBAL_REPORT_ACCESS_SCOPE,
        }
    ]

    captured.clear()
    assert client.get("/api/v1/reports/student-profile/xlsx").status_code == 200
    assert captured == [
        {
            "academic_year_id": None,
            "campus_id": None,
            "college_id": None,
            "program_id": None,
            "year_level": None,
            "access_scope": GLOBAL_REPORT_ACCESS_SCOPE,
        }
    ]


@pytest.mark.django_db
def test_xlsx_endpoint_passes_regular_counselor_college_scope(monkeypatch):
    sync_policy()
    counselor = make_user("scoped-xlsx@example.edu", "COUNSELOR")
    campus = Campus.objects.create(code="XLSX-SCOPE", name="XLSX Scope Campus")
    college = College.objects.create(campus=campus, code="XLSX-COL", name="XLSX Scope College")
    CounselorResponsibility.objects.create(college=college, counselor=counselor)
    captured: list[dict[str, object]] = []

    def fake_render(**kwargs):
        captured.append(kwargs)
        return fake_xlsx_result()

    monkeypatch.setattr(reports_api, "render_student_profiling_xlsx", fake_render)
    response = auth_client(counselor).get("/api/v1/reports/student-profile/xlsx")
    assert response.status_code == 200
    scope = captured[0]["access_scope"]
    assert scope.is_global is False
    assert scope.college_ids == (college.pk,)


@pytest.mark.django_db
@pytest.mark.parametrize(
    ("exc", "status_code", "error_code"),
    [
        (ReportNotFound("missing filter"), 404, "report_filter_not_found"),
        (
            ReportConfigurationConflict("configuration conflict"),
            409,
            "report_configuration_conflict",
        ),
        (InvalidReportFilter("contradictory filters"), 422, "invalid_report_filter"),
        (
            StudentProfilingWorkbookUnavailable("private workbook detail"),
            503,
            "report_workbook_unavailable",
        ),
    ],
)
def test_xlsx_endpoint_preserves_report_errors_and_hides_workbook_details(
    monkeypatch,
    exc,
    status_code,
    error_code,
):
    sync_policy()
    head = make_head(f"head-xlsx-error-{status_code}@example.edu")

    def fail(**kwargs):
        raise exc

    monkeypatch.setattr(reports_api, "render_student_profiling_xlsx", fail)
    response = auth_client(head).get("/api/v1/reports/student-profile/xlsx")
    assert response.status_code == status_code
    assert response.json()["error"]["code"] == error_code
    if status_code == 503:
        assert "private workbook detail" not in response.content.decode()


def test_render_service_calls_canonical_builder_once_and_forwards_exact_filters(monkeypatch):
    report = synthetic_report(program_count=3, submitted_count=3)
    calls: list[dict[str, object]] = []
    filters = {
        "academic_year_id": uuid4(),
        "campus_id": uuid4(),
        "college_id": uuid4(),
        "program_id": uuid4(),
        "year_level": 5,
    }

    def fake_builder(**kwargs):
        calls.append(kwargs)
        return report

    monkeypatch.setattr(report_xlsx, "build_student_profiling_report", fake_builder)
    scope = ReportAccessScope(is_global=False, college_ids=(uuid4(),))
    result = render_student_profiling_xlsx(**filters, access_scope=scope)

    assert calls == [{**filters, "access_scope": scope}]
    assert result.filename == "student-profile-2026-2027.xlsx"
    workbook = load_workbook(BytesIO(result.xlsx_bytes), data_only=False, keep_links=False)
    assert workbook.sheetnames == EXPECTED_SHEETS


def test_canonical_report_errors_remain_outside_workbook_error_boundary(monkeypatch):
    expected = ReportNotFound("canonical filter failure")

    def fail_builder(**kwargs):
        raise expected

    monkeypatch.setattr(report_xlsx, "build_student_profiling_report", fail_builder)
    with pytest.raises(ReportNotFound) as raised:
        render_student_profiling_xlsx(access_scope=GLOBAL_REPORT_ACCESS_SCOPE)
    assert raised.value is expected


def test_workbook_generation_failures_are_wrapped_without_internal_details(monkeypatch):
    monkeypatch.setattr(
        report_xlsx,
        "build_student_profiling_report",
        lambda **kwargs: synthetic_report(program_count=2),
    )

    def fail_workbook(report):
        raise ValueError("/private/path/openpyxl/zip detail")

    monkeypatch.setattr(report_xlsx, "build_student_profiling_workbook", fail_workbook)
    with pytest.raises(
        StudentProfilingWorkbookUnavailable,
        match="temporarily unavailable",
    ) as raised:
        render_student_profiling_xlsx(access_scope=GLOBAL_REPORT_ACCESS_SCOPE)
    assert "/private/path" not in str(raised.value)


def test_workbook_is_valid_deterministic_visible_and_formula_free(monkeypatch):
    report = synthetic_report(program_count=3, submitted_count=3)
    workbook = workbook_from_report(monkeypatch, report)

    assert workbook.sheetnames == EXPECTED_SHEETS
    assert len(workbook.worksheets) == 14
    assert getattr(workbook, "vba_archive", None) is None
    assert all(worksheet.sheet_state == "visible" for worksheet in workbook.worksheets)
    assert_no_formulas_or_hyperlinks(workbook)

    for _, sheet_name, _ in SECTION_SPECS:
        worksheet = workbook[sheet_name]
        assert worksheet.freeze_panes == "A2"
        assert worksheet.auto_filter.ref
        assert all(cell.font.bold for cell in worksheet[1])


def test_summary_preserves_timezone_context_current_coverage_and_methodology(monkeypatch):
    report = synthetic_report(program_count=2, submitted_count=2, mode="CURRENT")
    report["report_context"]["year_level_label"] = "3rd Year"
    report["inventory_coverage"]["ignored_filters"] = ["program_id", "year_level"]
    report["methodology"]["profile_population_note"] = "PROFILE POPULATION NOTE. " * 20
    # 07:00 UTC is 15:00 institutional time; the offset stays explicit.
    expected_timestamp = "2026-10-11T15:00:00+08:00"
    workbook = workbook_from_report(monkeypatch, report)
    values = workbook_values(workbook)

    for expected in (
        "Student Profile Report",
        "Individual Inventory-Based Student Profiling",
        "Academic Year",
        "2026-2027",
        "All Campuses",
        "All Colleges",
        "All Programs",
        "3rd Year",
        "Reporting Scope",
        "Institution-wide",
        "Generated At",
        expected_timestamp,
        "Submitted Inventory Profile Population (n)",
        "Current Student Inventory Coverage",
        "CURRENT",
        "Eligible Students",
        "Submitted",
        "Draft",
        "Without Individual Inventory",
        "Reconciliation of Coverage and Profile Population",
        "Submitted in coverage but outside the profile population",
        "CANONICAL CURRENT COVERAGE NOTE.",
        "Coverage filters not applicable to missing-Inventory classification: Program, Year Level",
        "Program Legend / Program Columns",
    ):
        assert expected in values

    # The reporting scope is stated once, not appended to every methodology note.
    assert values.count("Institution-wide") == 1
    assert "ACCESS SCOPE NOTE SHOULD NOT PRINT." not in values
    assert "CANONICAL CURRENT COVERAGE SCOPE NOTE." not in values

    summary = workbook["Summary"]
    methodology_cell = next(
        cell
        for row in summary.iter_rows()
        for cell in row
        if isinstance(cell.value, str) and cell.value.startswith("PROFILE POPULATION NOTE.")
    )
    assert methodology_cell.alignment.wrap_text is True
    # A long merged note gets an explicit height so it is readable without resizing rows.
    assert summary.row_dimensions[methodology_cell.row].height > 15


def test_historical_summary_does_not_fabricate_eligible_or_missing_counts(monkeypatch):
    report = synthetic_report(program_count=2, submitted_count=2, mode="HISTORICAL_LIMITED")
    workbook = workbook_from_report(monkeypatch, report)
    values = workbook_values(workbook)

    assert "HISTORICAL_LIMITED" in values
    assert "Inventory Coverage — Historical Academic Year (limited)" in values
    assert "Eligible Students" not in values
    assert "Not available from current COMPASS data" in values
    assert "CANONICAL HISTORICAL METHODOLOGY NOTE." in values
    assert "CANONICAL CURRENT COVERAGE NOTE." not in values


def test_section_sheet_uses_all_programs_key_lookup_and_canonical_percentage(monkeypatch):
    report = synthetic_report(program_count=4, submitted_count=4)
    programs = report["program_columns"]
    row = report["sections"]["sex"]["rows"][0]
    expected_counts = {
        str(program["key"]): (index + 1) * 10 for index, program in enumerate(programs)
    }
    row["program_counts"] = list(
        reversed([{"program_key": key, "count": count} for key, count in expected_counts.items()])
    )
    row["total_count"] = 777
    row["percentage"] = Decimal("32.72")

    workbook = workbook_from_report(monkeypatch, report)
    worksheet = workbook["Sex"]
    headers = [cell.value for cell in worksheet[1]]
    exported = [cell.value for cell in worksheet[2]]

    assert headers == [
        "Category",
        "CCMS-P01",
        "CCMS-P02",
        "CCMS-P03",
        "CCMS-P04",
        "Total",
        "Percentage (%)",
    ]
    assert exported[1:5] == [expected_counts[str(program["key"])] for program in programs]
    assert exported[5] == 777
    assert Decimal(str(exported[6])).quantize(Decimal("0.01")) == Decimal("32.72")
    assert worksheet.cell(row=2, column=7).number_format == "0.00"


def test_section_sheets_end_with_population_total_outside_the_filter(monkeypatch):
    report = synthetic_report(colleges=[("CCMS", 3)], submitted_count=11, legacy_inventories=2)
    workbook = workbook_from_report(monkeypatch, report)
    programs = report["program_columns"]

    for section_key, sheet_name, _ in SECTION_SPECS:
        worksheet = workbook[sheet_name]
        data_rows = len(report["sections"][section_key]["rows"])
        assert worksheet.auto_filter.ref.endswith(str(1 + data_rows))
        total_row = data_rows + 3
        offset = 3 if section_key == "city_municipality" else 1
        values = [cell.value for cell in worksheet[total_row]]
        assert values[0] == "Total (profile population)"
        assert values[offset : offset + len(programs)] == [
            program["submitted_inventory_count"] for program in programs
        ]
        assert values[offset + len(programs)] == 11
        assert Decimal(str(values[offset + len(programs) + 1])) == Decimal("100.00")
        assert str(worksheet.cell(row=total_row + 1, column=1).value).startswith(
            "Denominator: n = 11 submitted Individual Inventories."
        )
        # Every exported Program column sums to its canonical Program population.
        for index, program in enumerate(programs):
            column = offset + index + 1
            column_sum = sum(
                worksheet.cell(row=row, column=column).value for row in range(2, 2 + data_rows)
            )
            assert column_sum == program["submitted_inventory_count"]

    income = workbook["Parent Income"]
    income_values = [cell.value for row in income.iter_rows() for cell in row]
    assert "Parent Annual Income Brackets" in income_values
    assert "Below ₱131,484" in income_values
    assert "Above ₱2,626,680" in income_values


def test_duplicate_program_codes_are_disambiguated_and_legend_keeps_identity(monkeypatch):
    report = synthetic_report(program_count=2, submitted_count=2)
    first, second = report["program_columns"]
    first["code"] = second["code"] = "DUP"
    first["college"] = {"id": uuid4(), "code": "COL-A", "name": "College A"}
    first["campus"] = {"id": uuid4(), "code": "CAMP-A", "name": "Campus A"}
    second["college"] = {"id": uuid4(), "code": "COL-B", "name": "College B"}
    second["campus"] = {"id": uuid4(), "code": "CAMP-B", "name": "Campus B"}

    workbook = workbook_from_report(monkeypatch, report)
    sex_headers = [cell.value for cell in workbook["Sex"][1]]
    values = workbook_values(workbook)

    assert sex_headers[1] == "DUP (COL-A / CAMP-A)"
    assert sex_headers[2] == "DUP (COL-B / CAMP-B)"
    assert sex_headers[1] != sex_headers[2]
    assert "COL-A — College A" in values
    assert "COL-B — College B" in values
    assert first["name"] in values
    assert second["name"] in values


def test_legacy_program_column_is_retained(monkeypatch):
    report = synthetic_report(program_count=2, submitted_count=6, legacy_inventories=2)
    workbook = workbook_from_report(monkeypatch, report)
    assert "Not recorded / legacy" in [cell.value for cell in workbook["Sex"][1]]
    summary = workbook["Summary"]
    legend_rows = [
        [cell.value for cell in row]
        for row in summary.iter_rows()
        if row[0].value == "Not recorded / legacy"
    ]
    assert legend_rows[0][5:7] == ["Yes", 2]


def test_geography_sheet_preserves_canonical_city_province_region_context(monkeypatch):
    report = synthetic_report(program_count=2, submitted_count=2)
    base_row = report["sections"]["city_municipality"]["rows"][0]
    report["sections"]["city_municipality"]["rows"] = [
        base_row,
        {
            **base_row,
            "key": "NOT_SPECIFIED",
            "label": "Not specified",
            "province_name": None,
            "region_name": None,
        },
        {
            **base_row,
            "key": "NOT_RECORDED_LEGACY",
            "label": "Not recorded / legacy",
            "province_name": None,
            "region_name": None,
        },
    ]

    workbook = workbook_from_report(monkeypatch, report)
    worksheet = workbook["City Municipality"]
    assert [cell.value for cell in worksheet[1]][:3] == [
        "City / Municipality",
        "Province",
        "Region",
    ]
    assert [cell.value for cell in worksheet[2]][:3] == [
        "Synthetic Municipality 01",
        "Synthetic Province",
        "Synthetic Region",
    ]
    labels = [worksheet.cell(row=row, column=1).value for row in range(2, worksheet.max_row + 1)]
    assert "Not specified" in labels
    assert "Not recorded / legacy" in labels


def test_empty_report_keeps_stable_sheet_structure_with_headers_only(monkeypatch):
    report = synthetic_report(program_count=0, submitted_count=0)
    workbook = workbook_from_report(monkeypatch, report)

    assert workbook.sheetnames == EXPECTED_SHEETS
    assert EMPTY_REPORT_MESSAGE in workbook_values(workbook)
    for _, sheet_name, _ in SECTION_SPECS:
        worksheet = workbook[sheet_name]
        assert worksheet.max_row == 1
        assert worksheet.freeze_panes == "A2"


def test_untrusted_display_text_is_forced_to_text_cells_and_never_formulas(monkeypatch):
    report = synthetic_report(program_count=1, submitted_count=1)
    program = report["program_columns"][0]
    program["code"] = "=SUM(1,1)"
    program["name"] = "+1+1"
    program["college"] = {"id": uuid4(), "code": "-1+1", "name": ""}
    program["campus"] = {"id": uuid4(), "code": "@SUM(A1:A2)", "name": ""}
    report["sections"]["city_municipality"]["rows"][0]["label"] = "=CITY()"

    workbook = workbook_from_report(monkeypatch, report)
    dangerous = {"=SUM(1,1)", "+1+1", "-1+1", "@SUM(A1:A2)", "=CITY()"}
    seen: dict[str, str] = {}

    for worksheet in workbook.worksheets:
        for row in worksheet.iter_rows():
            for cell in row:
                if cell.value in dangerous:
                    seen[str(cell.value)] = cell.data_type
                assert cell.data_type != "f"
                assert cell.hyperlink is None

    assert seen == {value: "s" for value in dangerous}


def test_export_ignores_unexpected_private_keys_and_has_no_hidden_raw_sheet(monkeypatch):
    report = synthetic_report(program_count=2, submitted_count=2)
    secrets = {
        "PRIVATE-STUDENT-UUID",
        "PRIVATE-STUDENT-NUMBER",
        "private-student@example.edu",
        "PRIVATE-PHONE",
        "PRIVATE-EXACT-ADDRESS",
        "PRIVATE-PARENT-NAME",
        "PRIVATE-NARRATIVE",
    }
    report["private_student_uuid"] = "PRIVATE-STUDENT-UUID"
    report["private_student_number"] = "PRIVATE-STUDENT-NUMBER"
    report["private_email"] = "private-student@example.edu"
    report["private_phone"] = "PRIVATE-PHONE"
    report["private_address"] = "PRIVATE-EXACT-ADDRESS"
    report["private_parent_name"] = "PRIVATE-PARENT-NAME"
    report["sections"]["sex"]["rows"][0]["private_narrative"] = "PRIVATE-NARRATIVE"

    workbook = workbook_from_report(monkeypatch, report)
    serialized_values = {str(value) for value in workbook_values(workbook)}
    assert secrets.isdisjoint(serialized_values)
    assert set(workbook.sheetnames) == set(EXPECTED_SHEETS)
    assert all(worksheet.sheet_state == "visible" for worksheet in workbook.worksheets)


def test_filename_sanitization_is_shared_with_pdf_without_behavior_regression():
    assert student_profiling_xlsx_filename("2026-2027") == "student-profile-2026-2027.xlsx"
    assert student_profiling_xlsx_filename(" AY 2026/2027 ") == (
        "student-profile-AY-2026-2027.xlsx"
    )
    assert student_profiling_pdf_filename("2026-2027") == "student-profile-2026-2027.pdf"
    assert student_profiling_pdf_filename(" AY 2026/2027 ") == ("student-profile-AY-2026-2027.pdf")
    assert "/" not in student_profiling_xlsx_filename("../../AY 2026/2027")


def test_excel_column_limit_fails_safely_without_chunking_or_truncation():
    too_many_programs = [{}] * (MAX_EXCEL_COLUMNS - 4)
    with pytest.raises(
        StudentProfilingWorkbookUnavailable,
        match="column limit",
    ):
        _validate_excel_column_limit(too_many_programs)


def test_xlsx_includes_privacy_warning_and_na_without_suppressing_data(monkeypatch):
    report = synthetic_report(program_count=1, submitted_count=1)
    report["disclosure_warnings"] = [
        {
            "code": "SMALL_POPULATION",
            "message": (
                "This report includes results from a small population or category. "
                "Use and share these results carefully because some aggregate results "
                "may be easier to associate with individual respondents."
            ),
        }
    ]
    report["sections"]["sex"]["rows"][0]["percentage"] = None

    workbook = workbook_from_report(monkeypatch, report)
    values = workbook_values(workbook)
    sex = workbook["Sex"]

    assert "Small-population privacy notice" in values
    assert any(
        isinstance(value, str) and "Use and share these results carefully" in value
        for value in values
    )
    assert sex.cell(row=2, column=1).value == "Male"
    assert sex.cell(row=2, column=3).value == 1
    assert sex.cell(row=2, column=4).value == "N/A"


@pytest.mark.django_db
def test_pdf_and_xlsx_present_identical_canonical_numbers_from_one_frozen_report():
    sync_policy()
    year = AcademicYear.objects.create(label="2026-2027", is_current=True)
    revision = make_revision("parity")
    _, college_a, program_a1 = make_organization("PAR-A")
    program_a2 = Program.objects.create(college=college_a, code="P-PAR-A2", name="Program A2")
    _, _, program_b = make_organization("PAR-B")
    plan = [
        (program_a1, Sex.MALE, CivilStatusCategory.SINGLE),
        (program_a1, Sex.FEMALE, CivilStatusCategory.SINGLE),
        (program_a1, Sex.FEMALE, CivilStatusCategory.NOT_SPECIFIED),
        (program_a2, Sex.MALE, CivilStatusCategory.MARRIED),
        (program_a2, Sex.MALE, CivilStatusCategory.SINGLE),
        (program_b, Sex.FEMALE, CivilStatusCategory.SINGLE),
        (program_b, Sex.MALE, None),
        (None, Sex.FEMALE, None),
    ]
    for index, (program, sex, civil_status) in enumerate(plan):
        make_inventory(
            student=make_profile_student(f"parity-{index}@example.edu"),
            academic_year=year,
            revision=revision,
            program=program,
            sex=sex,
            civil_status_category=civil_status,
            date_of_birth=date(2007, 1, 1 + index),
            with_location=index % 2 == 0,
        )

    report = build_student_profiling_report(access_scope=GLOBAL_REPORT_ACCESS_SCOPE)
    frozen = copy.deepcopy(report)
    profile = build_student_profiling_print_context(report)["student_profile"]
    workbook = build_student_profiling_workbook(report)
    assert report == frozen, "presentation must never mutate the canonical report"

    pdf_tables = {table["key"]: table for part in profile["parts"] for table in part["tables"]}
    columns = report["program_columns"]
    for section_key, sheet_name, _ in SECTION_SPECS:
        worksheet = workbook[sheet_name]
        offset = 3 if section_key == "city_municipality" else 1
        total_column = offset + len(columns) + 1
        xlsx_rows = {}
        for row in range(2, 2 + len(report["sections"][section_key]["rows"])):
            label = worksheet.cell(row=row, column=1).value
            count = worksheet.cell(row=row, column=total_column).value
            percentage = worksheet.cell(row=row, column=total_column + 1).value
            xlsx_rows[label] = (count, f"{Decimal(str(percentage)):.2f}%")
        pdf_rows = {}
        for row in pdf_tables[section_key]["rows"][:-1]:
            cells = [cell["text"] for cell in row["cells"]]
            pdf_rows[cells[0]] = (int(cells[-2]), cells[-1])
        assert all(xlsx_rows[label] == value for label, value in pdf_rows.items()), section_key
        # Only zero-count Not recorded / legacy rows are left out of the printed table.
        for label in set(xlsx_rows) - set(pdf_rows):
            assert label == "Not recorded / legacy" and xlsx_rows[label][0] == 0
        footer = [cell.value for cell in worksheet[len(xlsx_rows) + 3]]
        pdf_total = [cell["text"] for cell in pdf_tables[section_key]["rows"][-1]["cells"]]
        assert footer[total_column - 1] == int(pdf_total[-2]) == len(plan)

    # College cross-tabs reconcile to the workbook's Program columns.
    sex_sheet = workbook["Sex"]
    headers = [cell.value for cell in sex_sheet[1]]
    for group in profile["appendix"]:
        table = group["tables"][0]
        program_headers = [
            column["label"]
            for column in table["columns"][1:]
            if column["label"] in headers and column["label"] != "Total"
        ]
        for row in table["rows"][:-1]:
            cells = [cell["text"] for cell in row["cells"]]
            sheet_row = next(
                index
                for index in range(2, sex_sheet.max_row + 1)
                if sex_sheet.cell(row=index, column=1).value == cells[0]
            )
            expected = [
                sex_sheet.cell(row=sheet_row, column=headers.index(label) + 1).value
                for label in program_headers
            ]
            assert [int(value) for value in cells[1 : 1 + len(program_headers)]] == expected
            if len(program_headers) > 1:
                assert int(cells[1 + len(program_headers)]) == sum(expected)
