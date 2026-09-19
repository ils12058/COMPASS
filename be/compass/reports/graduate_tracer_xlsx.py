"""Aggregate-only XLSX representation for Graduate Tracer reporting."""

from __future__ import annotations

from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal
from io import BytesIO

from openpyxl import Workbook
from openpyxl.cell.cell import Cell
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.worksheet import Worksheet

from .filenames import safe_report_filename_part
from .graduate_tracer import (
    GraduateTracerReportError,
    build_graduate_tracer_report,
    graduate_tracer_release_context,
)

XLSX_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
EMPTY_REPORT_MESSAGE = (
    "No submitted Graduate Tracer responses matched the selected submission period."
)

_SHEET_SECTIONS = {
    "Respondent Profile": (
        "sex",
        "civil_status",
        "region_of_origin",
        "residence_location",
    ),
    "Employment": (
        "current_employment_state",
        "present_employment_status",
        "employer_business_line",
        "place_of_work",
    ),
    "First Job": (
        "first_job_after_college",
        "first_job_related_to_course",
        "first_job_duration",
        "first_job_source",
        "time_to_first_job",
        "first_job_level",
        "current_job_level",
        "initial_gross_monthly_earning",
        "curriculum_relevant_to_first_job",
    ),
    "Reasons & Skills": (
        "unemployment_reasons",
        "reasons_for_staying_on_job",
        "useful_competencies",
    ),
}

_THIN = Side(style="thin", color="D9D9D9")
_TABLE_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
_HEADER_FILL = PatternFill(fill_type="solid", fgColor="E7E6E6")
_BOLD = Font(bold=True)
_WRAP = Alignment(wrap_text=True, vertical="top")
_CENTER = Alignment(horizontal="center", vertical="center")
_RIGHT = Alignment(horizontal="right", vertical="center")


class GraduateTracerWorkbookUnavailable(GraduateTracerReportError):
    """The aggregate is valid, but its XLSX representation is unavailable."""


@dataclass(frozen=True, slots=True)
class GraduateTracerXlsxResult:
    xlsx_bytes: bytes
    filename: str
    release_context: dict[str, object]


def _write_text(
    worksheet: Worksheet,
    row: int,
    column: int,
    value: object,
    *,
    bold: bool = False,
    wrap: bool = False,
) -> Cell:
    cell = worksheet.cell(row=row, column=column)
    cell.value = "" if value is None else str(value)
    cell.data_type = "s"
    cell.hyperlink = None
    if bold:
        cell.font = _BOLD
    if wrap:
        cell.alignment = _WRAP
    return cell


def _write_number(
    worksheet: Worksheet,
    row: int,
    column: int,
    value: int | Decimal,
    *,
    number_format: str | None = None,
) -> Cell:
    cell = worksheet.cell(row=row, column=column, value=value)
    cell.hyperlink = None
    cell.alignment = _RIGHT
    if number_format is not None:
        cell.number_format = number_format
    return cell


def _as_decimal(value: object) -> Decimal:
    if isinstance(value, Decimal):
        return value
    if isinstance(value, int | float):
        return Decimal(str(value))
    raise GraduateTracerWorkbookUnavailable(
        "A Graduate Tracer percentage value is invalid."
    )


def _date_text(value: object) -> str:
    if value is None:
        return "All"
    if isinstance(value, date):
        return value.isoformat()
    raise GraduateTracerWorkbookUnavailable(
        "A Graduate Tracer submission-period value is invalid."
    )


def _generated_at_text(value: object) -> str:
    if not isinstance(value, datetime) or value.utcoffset() is None:
        raise GraduateTracerWorkbookUnavailable(
            "The Graduate Tracer generated timestamp is invalid."
        )
    return value.isoformat(timespec="seconds")


def _build_summary(
    worksheet: Worksheet,
    *,
    report_context: dict[str, object],
    methodology: dict[str, object],
) -> None:
    worksheet.title = "Summary"
    worksheet.sheet_view.showGridLines = False
    _write_text(worksheet, 1, 1, "Graduate Tracer Aggregate Report", bold=True)
    worksheet["A1"].font = Font(bold=True, size=14)

    rows = [
        ("Instrument Schema Version", report_context.get("instrument_schema_version")),
        ("Submitted From", _date_text(report_context.get("submitted_from"))),
        ("Submitted To", _date_text(report_context.get("submitted_to"))),
        ("Submitted Response Count", int(report_context.get("submitted_response_count") or 0)),
        ("Generated At", _generated_at_text(report_context.get("generated_at"))),
    ]
    row = 3
    for label, value in rows:
        _write_text(worksheet, row, 1, label, bold=True, wrap=True)
        if label == "Submitted Response Count":
            _write_number(worksheet, row, 2, int(value))
        else:
            _write_text(worksheet, row, 2, value, wrap=True)
        row += 1

    if int(report_context.get("submitted_response_count") or 0) == 0:
        row += 1
        _write_text(worksheet, row, 1, EMPTY_REPORT_MESSAGE, bold=True, wrap=True)
        row += 2

    _write_text(worksheet, row, 1, "Methodology and Limitations", bold=True)
    row += 1
    for value in methodology.values():
        _write_text(worksheet, row, 1, value, wrap=True)
        row += 1

    worksheet.column_dimensions["A"].width = 44
    worksheet.column_dimensions["B"].width = 36


def _write_section(
    worksheet: Worksheet,
    *,
    start_row: int,
    section: dict[str, object],
) -> int:
    label = str(section.get("label") or "")
    denominator = int(section.get("denominator") or 0)
    denominator_label = str(section.get("denominator_label") or "")
    multiple_selection = bool(section.get("multiple_selection"))
    rows = section.get("rows")
    if not isinstance(rows, list):
        raise GraduateTracerWorkbookUnavailable(
            "A Graduate Tracer distribution section is invalid."
        )

    _write_text(worksheet, start_row, 1, label, bold=True, wrap=True)
    _write_text(worksheet, start_row + 1, 1, "Denominator", bold=True)
    _write_number(worksheet, start_row + 1, 2, denominator)
    _write_text(worksheet, start_row + 2, 1, "Applicable Population", bold=True)
    _write_text(worksheet, start_row + 2, 2, denominator_label, wrap=True)
    _write_text(worksheet, start_row + 3, 1, "Multiple Selection", bold=True)
    _write_text(
        worksheet,
        start_row + 3,
        2,
        "Yes" if multiple_selection else "No",
    )

    header_row = start_row + 5
    for column, header in enumerate(("Category", "Count", "Percentage (%)"), start=1):
        cell = _write_text(worksheet, header_row, column, header, bold=True, wrap=True)
        cell.fill = _HEADER_FILL
        cell.border = _TABLE_BORDER
        cell.alignment = _CENTER

    row_number = header_row + 1
    for raw in rows:
        if not isinstance(raw, dict):
            raise GraduateTracerWorkbookUnavailable(
                "A Graduate Tracer distribution row is invalid."
            )
        text_cell = _write_text(
            worksheet,
            row_number,
            1,
            raw.get("label"),
            wrap=True,
        )
        text_cell.border = _TABLE_BORDER

        count_cell = _write_number(
            worksheet,
            row_number,
            2,
            int(raw.get("count") or 0),
        )
        count_cell.border = _TABLE_BORDER

        percentage_cell = _write_number(
            worksheet,
            row_number,
            3,
            _as_decimal(raw.get("percentage")),
            number_format="0.00",
        )
        percentage_cell.border = _TABLE_BORDER
        row_number += 1
    return row_number + 2


def _build_section_sheet(
    workbook: Workbook,
    *,
    title: str,
    section_keys: tuple[str, ...],
    sections: dict[str, object],
) -> None:
    worksheet = workbook.create_sheet(title=title)
    worksheet.sheet_view.showGridLines = False
    worksheet.sheet_state = "visible"
    row = 1
    for key in section_keys:
        raw = sections.get(key)
        if not isinstance(raw, dict):
            raise GraduateTracerWorkbookUnavailable(
                f"The Graduate Tracer {key} section is unavailable."
            )
        row = _write_section(worksheet, start_row=row, section=raw)

    worksheet.column_dimensions["A"].width = 48
    worksheet.column_dimensions["B"].width = 28
    worksheet.column_dimensions["C"].width = 18


def _filename(*, submitted_from: date | None, submitted_to: date | None) -> str:
    stem = "graduate-tracer-schema-v1"
    if submitted_from is not None or submitted_to is not None:
        start = submitted_from.isoformat() if submitted_from else "start"
        end = submitted_to.isoformat() if submitted_to else "end"
        stem += f"-submitted-{start}-to-{end}"
    return f"{safe_report_filename_part(stem, fallback='graduate-tracer-schema-v1')}.xlsx"


def render_graduate_tracer_xlsx(
    *,
    submitted_from: date | None = None,
    submitted_to: date | None = None,
) -> GraduateTracerXlsxResult:
    report = build_graduate_tracer_report(
        submitted_from=submitted_from,
        submitted_to=submitted_to,
    )
    report_context = report.get("report_context")
    methodology = report.get("methodology")
    sections = report.get("sections")
    if not isinstance(report_context, dict):
        raise GraduateTracerWorkbookUnavailable(
            "The Graduate Tracer report context is unavailable."
        )
    if not isinstance(methodology, dict):
        raise GraduateTracerWorkbookUnavailable(
            "The Graduate Tracer methodology is unavailable."
        )
    if not isinstance(sections, dict):
        raise GraduateTracerWorkbookUnavailable(
            "The Graduate Tracer report sections are unavailable."
        )

    try:
        workbook = Workbook()
        summary = workbook.active
        if summary is None:
            raise GraduateTracerWorkbookUnavailable(
                "The Graduate Tracer Summary worksheet is unavailable."
            )
        _build_summary(
            summary,
            report_context=report_context,
            methodology=methodology,
        )
        summary.sheet_state = "visible"
        for title, keys in _SHEET_SECTIONS.items():
            _build_section_sheet(
                workbook,
                title=title,
                section_keys=keys,
                sections=sections,
            )
        for worksheet in workbook.worksheets:
            worksheet.sheet_state = "visible"

        output = BytesIO()
        workbook.save(output)
    except GraduateTracerReportError:
        raise
    except Exception as exc:
        raise GraduateTracerWorkbookUnavailable(
            "The Graduate Tracer XLSX export could not be generated."
        ) from exc

    return GraduateTracerXlsxResult(
        xlsx_bytes=output.getvalue(),
        filename=_filename(
            submitted_from=submitted_from,
            submitted_to=submitted_to,
        ),
        release_context=graduate_tracer_release_context(
            submitted_from=submitted_from,
            submitted_to=submitted_to,
        ),
    )


__all__ = [
    "EMPTY_REPORT_MESSAGE",
    "GraduateTracerWorkbookUnavailable",
    "GraduateTracerXlsxResult",
    "XLSX_CONTENT_TYPE",
    "render_graduate_tracer_xlsx",
]
