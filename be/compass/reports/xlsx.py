"""Spreadsheet-native Student Profiling XLSX export over the canonical aggregate result."""

from __future__ import annotations

import textwrap
from dataclasses import dataclass
from datetime import datetime
from decimal import Decimal
from io import BytesIO
from uuid import UUID

from openpyxl import Workbook
from openpyxl.cell.cell import Cell
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from compass.common.institutional_time import to_institution_time

from .filenames import safe_report_filename_part
from .services import (
    PARENT_INCOME_LEGEND_NOTE,
    ReportAccessScope,
    ReportError,
    build_student_profiling_report,
    parent_income_bracket_legend,
    student_profiling_release_context,
)

XLSX_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
MAX_EXCEL_COLUMNS = 16_384
EMPTY_REPORT_MESSAGE = "No submitted Individual Inventories matched the selected profile filters."

FILTER_LABELS = {
    "academic_year_id": "Academic Year",
    "campus_id": "Campus",
    "college_id": "College",
    "program_id": "Program",
    "year_level": "Year Level",
}

SECTION_SPECS = (
    ("sex", "Sex", "Category"),
    ("age", "Age", "Age"),
    ("civil_status", "Civil Status", "Category"),
    ("physical_disadvantage", "PWD Status", "Category"),
    ("current_religion", "Religion", "Category"),
    ("mother_life_status", "Mother Life", "Category"),
    ("father_life_status", "Father Life", "Category"),
    ("parent_family_status", "Parent Family", "Category"),
    ("city_municipality", "City Municipality", "City / Municipality"),
    ("parent_annual_income", "Parent Income", "Category"),
    ("mother_occupation", "Mother Occupation", "Category"),
    ("father_occupation", "Father Occupation", "Category"),
    ("living_condition", "Living Condition", "Category"),
)

_THIN_SIDE = Side(style="thin", color="B7B7B7")
_TABLE_BORDER = Border(
    left=_THIN_SIDE,
    right=_THIN_SIDE,
    top=_THIN_SIDE,
    bottom=_THIN_SIDE,
)
_HEADER_FILL = PatternFill(fill_type="solid", fgColor="E7E6E6")
_TOTAL_FILL = PatternFill(fill_type="solid", fgColor="F2F2F2")
_BOLD = Font(bold=True)
_WRAP = Alignment(wrap_text=True, vertical="top")
# Summary notes span columns A:G. Merged cells do not grow automatically, so note rows get an
# explicit height from their length: about this many characters fit on one line of the merge.
_NOTE_CHARACTERS_PER_LINE = 120
_NOTE_LINE_HEIGHT = 15
SUMMARY_WIDTH = 7
_CENTER = Alignment(horizontal="center", vertical="center")
_RIGHT = Alignment(horizontal="right", vertical="center")


class StudentProfilingWorkbookUnavailable(ReportError):
    """The canonical aggregate is valid, but its XLSX representation is unavailable."""


@dataclass(frozen=True, slots=True)
class StudentProfilingXlsxResult:
    xlsx_bytes: bytes
    filename: str
    release_context: dict[str, object]


def _require_dict(value: object, label: str) -> dict[str, object]:
    if not isinstance(value, dict):
        raise StudentProfilingWorkbookUnavailable(
            f"The Student Profiling {label} is unavailable for XLSX export."
        )
    return value


def _require_list(value: object, label: str) -> list[object]:
    if not isinstance(value, list):
        raise StudentProfilingWorkbookUnavailable(
            f"The Student Profiling {label} is unavailable for XLSX export."
        )
    return value


def _write_text(
    worksheet: Worksheet,
    row: int,
    column: int,
    value: object,
    *,
    bold: bool = False,
    wrap: bool = False,
) -> Cell:
    """Write display text explicitly as a string cell, never as an Excel formula."""
    cell = worksheet.cell(row=row, column=column)
    cell.value = "" if value is None else str(value)
    cell.data_type = "s"
    cell.hyperlink = None
    if bold:
        cell.font = _BOLD
    if wrap:
        cell.alignment = _WRAP
    return cell


def _write_table_header(worksheet: Worksheet, row: int, headers: list[str]) -> None:
    for column, header in enumerate(headers, start=1):
        cell = _write_text(worksheet, row, column, header, bold=True, wrap=True)
        cell.fill = _HEADER_FILL
        cell.border = _TABLE_BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    worksheet.row_dimensions[row].height = 30


def _write_table_text(worksheet: Worksheet, row: int, column: int, value: object) -> Cell:
    cell = _write_text(worksheet, row, column, value, wrap=True)
    cell.border = _TABLE_BORDER
    return cell


def _write_table_number(
    worksheet: Worksheet,
    row: int,
    column: int,
    value: int | Decimal,
    *,
    number_format: str | None = None,
) -> Cell:
    cell = worksheet.cell(row=row, column=column, value=value)
    cell.border = _TABLE_BORDER
    cell.alignment = _RIGHT
    if number_format is not None:
        cell.number_format = number_format
    return cell


def _set_capped_width(
    worksheet: Worksheet,
    column: int,
    *,
    minimum: float,
    maximum: float,
) -> None:
    max_length = 0
    for cell in worksheet[get_column_letter(column)]:
        if cell.value is None:
            continue
        max_length = max(max_length, len(str(cell.value)))
    width = max(minimum, min(maximum, max_length + 2))
    worksheet.column_dimensions[get_column_letter(column)].width = width


def _organization_display(item: object, fallback: str) -> str:
    if item is None:
        return fallback
    value = _require_dict(item, "organization context")
    code = str(value.get("code") or "").strip()
    name = str(value.get("name") or "").strip()
    if code and name:
        return f"{code} — {name}"
    return code or name or fallback


def _program_display_columns(program_columns: object) -> list[dict[str, object]]:
    raw_columns = _require_list(program_columns, "Program columns")
    seen_keys: set[str] = set()
    code_counts: dict[str, int] = {}

    for raw_column in raw_columns:
        column = _require_dict(raw_column, "Program column")
        key = str(column.get("key") or "")
        if not key or key in seen_keys:
            raise StudentProfilingWorkbookUnavailable(
                "The Student Profiling Program-column identity is inconsistent."
            )
        seen_keys.add(key)
        code = str(column.get("code") or "").strip()
        if code:
            code_counts[code] = code_counts.get(code, 0) + 1

    display_columns: list[dict[str, object]] = []
    for raw_column in raw_columns:
        column = _require_dict(raw_column, "Program column")
        key = str(column["key"])
        code = str(column.get("code") or "").strip()
        name = str(column.get("name") or "").strip()
        college = column.get("college")
        campus = column.get("campus")
        is_legacy = bool(column.get("is_legacy"))

        college_dict = None if college is None else _require_dict(college, "Program College")
        campus_dict = None if campus is None else _require_dict(campus, "Program Campus")

        if is_legacy:
            display_label = name or "Not recorded / legacy"
        elif code:
            display_label = code
            if code_counts.get(code, 0) > 1:
                context_codes = [
                    str(parent.get("code") or "").strip()
                    for parent in (college_dict, campus_dict)
                    if parent is not None and str(parent.get("code") or "").strip()
                ]
                if context_codes:
                    display_label = f"{code} ({' / '.join(context_codes)})"
        else:
            display_label = name or "Program"

        submitted = column.get("submitted_inventory_count")
        if type(submitted) is not int or submitted < 0:
            raise StudentProfilingWorkbookUnavailable(
                "The Student Profiling Program population is invalid."
            )
        display_columns.append(
            {
                "key": key,
                "display_label": display_label,
                "submitted_inventory_count": submitted,
                "code": code,
                "name": name or ("Not recorded / legacy" if is_legacy else "Program"),
                "college": _organization_display(college_dict, "") if college_dict else "",
                "campus": _organization_display(campus_dict, "") if campus_dict else "",
                "is_legacy": is_legacy,
                "college_id": str(college_dict["id"]) if college_dict else None,
            }
        )
    return display_columns


def _program_count_lookup(
    row: dict[str, object],
    program_columns: list[dict[str, object]],
) -> dict[str, int]:
    raw_counts = _require_list(row.get("program_counts"), "Program counts")
    lookup: dict[str, int] = {}
    for raw_count in raw_counts:
        count = _require_dict(raw_count, "Program count")
        key = str(count.get("program_key") or "")
        if not key or key in lookup:
            raise StudentProfilingWorkbookUnavailable(
                "The Student Profiling Program counts contain an invalid or duplicate identity."
            )
        lookup[key] = int(count.get("count") or 0)

    expected = {str(column["key"]) for column in program_columns}
    if set(lookup) != expected:
        raise StudentProfilingWorkbookUnavailable(
            "The Student Profiling Program counts do not match the canonical Program columns."
        )
    return lookup


def _generated_at_text(value: object) -> str:
    if not isinstance(value, datetime) or value.utcoffset() is None:
        raise StudentProfilingWorkbookUnavailable(
            "The Student Profiling generated timestamp must include timezone information."
        )
    # Same institutional civil time as the PDF, kept as ISO-8601 with its offset.
    return to_institution_time(value).isoformat(timespec="seconds")


def _percentage_value(value: object) -> Decimal | None:
    if value is None:
        return None
    if isinstance(value, Decimal):
        return value
    if isinstance(value, int | float):
        return Decimal(str(value))
    raise StudentProfilingWorkbookUnavailable("The Student Profiling percentage value is invalid.")


def _coverage_heading(coverage: dict[str, object]) -> str:
    mode = str(coverage.get("mode") or "")
    if mode == "CURRENT":
        return "Current Student Inventory Coverage"
    if mode == "HISTORICAL_LIMITED":
        return "Inventory Coverage — Historical Academic Year (limited)"
    raise StudentProfilingWorkbookUnavailable("The Student Profiling coverage mode is unsupported.")


def _coverage_rows(coverage: dict[str, object]) -> list[tuple[str, object]]:
    mode = str(coverage.get("mode") or "")
    if mode == "CURRENT":
        return [
            ("Coverage Mode", mode),
            ("Eligible Students", coverage.get("eligible_student_count")),
            ("Submitted", coverage.get("submitted_count")),
            ("Draft", coverage.get("draft_count")),
            ("Without Individual Inventory", coverage.get("missing_count")),
        ]
    if mode == "HISTORICAL_LIMITED":
        return [
            ("Coverage Mode", mode),
            ("Submitted", coverage.get("submitted_count")),
            ("Draft", coverage.get("draft_count")),
            (
                "Without Individual Inventory",
                "Not available from current COMPASS data",
            ),
        ]
    raise StudentProfilingWorkbookUnavailable("The Student Profiling coverage mode is unsupported.")


def _reconciliation_rows(coverage: dict[str, object]) -> list[tuple[str, int]]:
    raw = _require_dict(coverage.get("reconciliation"), "population reconciliation")
    values: dict[str, int] = {}
    for key in (
        "profile_submitted_count",
        "coverage_submitted_count",
        "in_both_count",
        "profile_only_count",
        "coverage_only_count",
    ):
        value = raw.get(key)
        if type(value) is not int or value < 0:
            raise StudentProfilingWorkbookUnavailable(
                "The Student Profiling population reconciliation is invalid."
            )
        values[key] = value
    return [
        ("Submitted Inventories in the profile population", values["profile_submitted_count"]),
        ("Also counted as Submitted in coverage", values["in_both_count"]),
        ("Not counted as Submitted in coverage", values["profile_only_count"]),
        ("Counted as Submitted in coverage", values["coverage_submitted_count"]),
        ("Submitted in coverage but outside the profile population", values["coverage_only_count"]),
    ]


def _coverage_ignored_note(coverage: dict[str, object]) -> str | None:
    raw_filters = coverage.get("ignored_filters", [])
    if not isinstance(raw_filters, list):
        raise StudentProfilingWorkbookUnavailable(
            "The Student Profiling ignored coverage filters are invalid."
        )
    labels = [FILTER_LABELS.get(str(item), str(item)) for item in raw_filters]
    if not labels:
        return None
    return "Coverage filters not applicable to missing-Inventory classification: " + ", ".join(
        labels
    )


def _validate_excel_column_limit(program_columns: list[dict[str, object]]) -> None:
    widest_section = len(program_columns) + 5
    if widest_section > MAX_EXCEL_COLUMNS:
        raise StudentProfilingWorkbookUnavailable(
            "The Student Profiling report exceeds the XLSX worksheet column limit."
        )


def _write_note(worksheet: Worksheet, row: int, text: object, *, bold: bool = False) -> None:
    value = "" if text is None else str(text)
    _write_text(worksheet, row, 1, value, bold=bold, wrap=True)
    worksheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=SUMMARY_WIDTH)
    lines = max(1, -(-len(value) // _NOTE_CHARACTERS_PER_LINE))
    if lines > 1:
        worksheet.row_dimensions[row].height = lines * _NOTE_LINE_HEIGHT


def _summary_heading(worksheet: Worksheet, row: int, label: str) -> None:
    _write_text(worksheet, row, 1, label, bold=True)
    worksheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=SUMMARY_WIDTH)
    cell = worksheet.cell(row=row, column=1)
    cell.fill = _HEADER_FILL
    cell.border = _TABLE_BORDER
    cell.font = Font(bold=True, color="5B1F26")
    worksheet.row_dimensions[row].height = 21


def _set_summary_row_height(worksheet: Worksheet, row: int, value: object) -> None:
    width = max(
        1,
        sum(
            float(worksheet.column_dimensions[get_column_letter(column)].width or 10)
            for column in range(2, SUMMARY_WIDTH + 1)
        )
        - 2,
    )
    lines = max(1, len(textwrap.wrap(str(value or ""), width=width, break_long_words=True)))
    worksheet.row_dimensions[row].height = max(18, lines * _NOTE_LINE_HEIGHT)


def _set_wrapped_row_heights(worksheet: Worksheet, *, first_row: int, last_row: int) -> None:
    merged_ranges = list(worksheet.merged_cells.ranges)
    for row_number in range(first_row, last_row + 1):
        line_count = 1
        for cell in worksheet[row_number]:
            if not cell.alignment.wrap_text or cell.value is None:
                continue
            merged_range = next(
                (
                    area
                    for area in merged_ranges
                    if area.min_row == row_number
                    and area.min_col == cell.column
                    and area.max_col > area.min_col
                ),
                None,
            )
            if merged_range is not None:
                width = sum(
                    float(worksheet.column_dimensions[get_column_letter(column)].width or 10)
                    for column in range(merged_range.min_col, merged_range.max_col + 1)
                )
            else:
                dimension = worksheet.column_dimensions[get_column_letter(cell.column)]
                width = float(dimension.width or 10)
            width = max(1, int(width) - 2)
            line_count = max(
                line_count,
                len(textwrap.wrap(str(cell.value), width=width, break_long_words=True)),
            )
        if line_count > 1:
            worksheet.row_dimensions[row_number].height = max(18, line_count * 15)


def _set_summary_print_layout(worksheet: Worksheet) -> None:
    worksheet.page_setup.orientation = "landscape"
    worksheet.page_setup.paperSize = worksheet.PAPERSIZE_A4
    worksheet.page_setup.fitToWidth = 1
    worksheet.page_setup.fitToHeight = 0
    worksheet.sheet_properties.pageSetUpPr.fitToPage = True
    worksheet.page_margins.left = 0.35
    worksheet.page_margins.right = 0.35
    worksheet.page_margins.top = 0.5
    worksheet.page_margins.bottom = 0.5
    worksheet.print_area = f"A1:G{worksheet.max_row}"


def _reporting_scope_text(report_context: dict[str, object]) -> str:
    scope = _require_dict(report_context.get("reporting_scope"), "reporting scope")
    if scope.get("is_global") is True:
        return "Institution-wide"
    labels = []
    for raw in _require_list(scope.get("colleges", []), "reporting scope Colleges"):
        college = _require_dict(raw, "reporting scope College")
        labels.append(_organization_display(college, "College"))
    if not labels:
        raise StudentProfilingWorkbookUnavailable(
            "The Student Profiling reporting scope has no authorized College."
        )
    return "Limited to assigned Colleges: " + "; ".join(labels)


def _build_summary_sheet(
    worksheet: Worksheet,
    report_context: dict[str, object],
    methodology: dict[str, object],
    coverage: dict[str, object],
    disclosure_warnings: list[object],
    program_columns: list[dict[str, object]],
) -> None:
    worksheet.title = "Summary"
    worksheet.sheet_view.showGridLines = False

    _write_text(worksheet, 1, 1, "Student Profile Report", bold=True)
    worksheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=SUMMARY_WIDTH)
    worksheet["A1"].alignment = Alignment(horizontal="center", vertical="center")
    worksheet["A1"].font = Font(bold=True, size=14)
    _write_text(worksheet, 2, 1, "Individual Inventory-Based Student Profiling")
    worksheet.merge_cells(start_row=2, start_column=1, end_row=2, end_column=SUMMARY_WIDTH)
    worksheet["A2"].alignment = Alignment(horizontal="center", vertical="center")

    academic_year = _require_dict(report_context.get("academic_year"), "Academic Year context")
    context_rows = [
        ("Academic Year", academic_year.get("label")),
        ("Campus filter", _organization_display(report_context.get("campus"), "All Campuses")),
        ("College filter", _organization_display(report_context.get("college"), "All Colleges")),
        ("Program filter", _organization_display(report_context.get("program"), "All Programs")),
        ("Year Level filter", report_context.get("year_level_label") or "All Year Levels"),
        ("Reporting Scope", _reporting_scope_text(report_context)),
        ("Generated At", _generated_at_text(report_context.get("generated_at"))),
    ]

    _summary_heading(worksheet, 4, "Report Context")
    row = 5
    for label, value in context_rows:
        _write_text(worksheet, row, 1, label, bold=True, wrap=True)
        _write_text(worksheet, row, 2, value, wrap=True)
        worksheet.merge_cells(start_row=row, start_column=2, end_row=row, end_column=SUMMARY_WIDTH)
        _set_summary_row_height(worksheet, row, value)
        row += 1

    row += 1
    _summary_heading(worksheet, row, "Key Figures")
    row += 1
    population = int(report_context.get("submitted_inventory_count") or 0)
    _write_text(
        worksheet,
        row,
        1,
        "Submitted Inventory Profile Population (n)",
        bold=True,
        wrap=True,
    )
    _write_table_number(worksheet, row, 2, population, number_format="#,##0")
    worksheet.merge_cells(start_row=row, start_column=2, end_row=row, end_column=SUMMARY_WIDTH)
    row += 1
    _write_text(worksheet, row, 1, _coverage_heading(coverage), bold=True, wrap=True)
    row += 1
    for label, value in _coverage_rows(coverage):
        _write_text(worksheet, row, 1, label, bold=True, wrap=True)
        if isinstance(value, int):
            _write_table_number(worksheet, row, 2, value, number_format="#,##0")
        else:
            _write_text(worksheet, row, 2, value, wrap=True)
        worksheet.merge_cells(start_row=row, start_column=2, end_row=row, end_column=SUMMARY_WIDTH)
        if not isinstance(value, int):
            _set_summary_row_height(worksheet, row, value)
        row += 1

    if population > 0:
        college_ids = {
            str(program["college_id"])
            for program in program_columns
            if not program["is_legacy"] and program.get("college_id")
        }
        represented_programs = sum(not program["is_legacy"] for program in program_columns)
        _write_text(worksheet, row, 1, "Recorded Colleges represented", bold=True, wrap=True)
        _write_table_number(worksheet, row, 2, len(college_ids), number_format="#,##0")
        worksheet.merge_cells(start_row=row, start_column=2, end_row=row, end_column=SUMMARY_WIDTH)
        row += 1
        _write_text(worksheet, row, 1, "Recorded Programs represented", bold=True, wrap=True)
        _write_table_number(worksheet, row, 2, represented_programs, number_format="#,##0")
        worksheet.merge_cells(start_row=row, start_column=2, end_row=row, end_column=SUMMARY_WIDTH)
        row += 1

    warning_messages: list[str] = []
    for warning in disclosure_warnings:
        item = _require_dict(warning, "disclosure warning")
        message = str(item.get("message") or "").strip()
        if message and message not in warning_messages:
            warning_messages.append(message)
    if warning_messages:
        row += 1
        _summary_heading(worksheet, row, "Small-population privacy notice")
        for message in warning_messages:
            row += 1
            _write_note(worksheet, row, message)
            worksheet.cell(row=row, column=1).fill = PatternFill(
                fill_type="solid", fgColor="FFF8E7"
            )
        row += 1
        _write_note(worksheet, row, "Exact counts are exported; no value is suppressed.")

    if population == 0:
        row += 1
        _write_note(worksheet, row, EMPTY_REPORT_MESSAGE, bold=True)

    row += 2
    _summary_heading(worksheet, row, "Reconciliation of Coverage and Profile Population")
    row += 1
    for label, value in _reconciliation_rows(coverage):
        _write_text(worksheet, row, 1, label, bold=True, wrap=True)
        _write_table_number(worksheet, row, 2, value, number_format="#,##0")
        worksheet.merge_cells(start_row=row, start_column=2, end_row=row, end_column=SUMMARY_WIDTH)
        row += 1
    _write_note(
        worksheet,
        row,
        "The two populations are counted on different bases, so no submission rate is "
        "calculated from them. Profile Inventories outside coverage belong to Students who are no "
        "longer active Current Students or are currently affiliated outside the coverage scope; "
        "coverage Inventories outside the profile record a Program, College, Campus, or Year "
        "Level outside the selected profile filters, or no Program.",
    )

    row += 2
    _summary_heading(worksheet, row, "Program Legend / Program Columns")
    row += 1
    legend_headers = [
        "Displayed Column",
        "Program Code",
        "Program Name",
        "College",
        "Campus",
        "Legacy",
        "Submitted Inventories",
    ]
    _write_table_header(worksheet, row, legend_headers)
    for program in program_columns:
        row += 1
        values = [
            program["display_label"],
            program["code"],
            program["name"],
            program["college"],
            program["campus"],
            "Yes" if program["is_legacy"] else "No",
        ]
        for column, value in enumerate(values, start=1):
            _write_table_text(worksheet, row, column, value)
        _write_table_number(
            worksheet,
            row,
            7,
            int(program["submitted_inventory_count"]),
            number_format="#,##0",
        )

    row += 2
    _summary_heading(worksheet, row, "Methodology and Coverage Notes")
    row += 1
    _write_note(worksheet, row, methodology.get("profile_population_note"))
    row += 1
    coverage_note = (
        methodology.get("historical_coverage_note")
        if coverage.get("mode") == "HISTORICAL_LIMITED"
        else methodology.get("coverage_note")
    )
    _write_note(worksheet, row, coverage_note or coverage.get("scope_note"))
    ignored_note = _coverage_ignored_note(coverage)
    if ignored_note:
        row += 1
        _write_note(worksheet, row, ignored_note)
    row += 1
    _write_note(
        worksheet,
        row,
        "Category sheets use the profile population as the denominator. Percentages are rounded "
        "half-up to two decimals independently, so a column may add up to 99.99 or 100.01. Each "
        "sheet ends with a Total row giving every Program's submitted Inventories.",
    )

    for column, width in enumerate((36, 18, 22, 22, 22, 14, 24), start=1):
        worksheet.column_dimensions[get_column_letter(column)].width = width
    _set_wrapped_row_heights(worksheet, first_row=3, last_row=worksheet.max_row)
    worksheet.freeze_panes = "A4"
    _set_summary_print_layout(worksheet)


def _build_section_sheet(
    workbook: Workbook,
    *,
    section_key: str,
    sheet_name: str,
    category_header: str,
    section: dict[str, object],
    program_columns: list[dict[str, object]],
    empty_report: bool,
) -> None:
    worksheet = workbook.create_sheet(title=sheet_name)
    is_geography = section_key == "city_municipality"
    headers = [category_header, "Province", "Region"] if is_geography else [category_header]
    headers.extend(str(program["display_label"]) for program in program_columns)
    headers.extend(["Total", "Percentage (%)"])
    _write_table_header(worksheet, 1, headers)

    if not empty_report:
        raw_rows = _require_list(section.get("rows"), f"{section_key} rows")
        for row_number, raw_row in enumerate(raw_rows, start=2):
            row = _require_dict(raw_row, f"{section_key} row")
            count_lookup = _program_count_lookup(row, program_columns)
            column_number = 1
            _write_table_text(worksheet, row_number, column_number, row.get("label"))
            column_number += 1

            if is_geography:
                _write_table_text(
                    worksheet,
                    row_number,
                    column_number,
                    row.get("province_name") or "",
                )
                column_number += 1
                _write_table_text(
                    worksheet,
                    row_number,
                    column_number,
                    row.get("region_name") or "",
                )
                column_number += 1

            for program in program_columns:
                _write_table_number(
                    worksheet,
                    row_number,
                    column_number,
                    count_lookup[str(program["key"])],
                )
                column_number += 1

            _write_table_number(
                worksheet,
                row_number,
                column_number,
                int(row.get("total_count") or 0),
            )
            column_number += 1
            percentage = _percentage_value(row.get("percentage"))
            if percentage is None:
                _write_table_text(
                    worksheet,
                    row_number,
                    column_number,
                    "N/A",
                )
            else:
                _write_table_number(
                    worksheet,
                    row_number,
                    column_number,
                    percentage,
                    number_format="0.00",
                )

    worksheet.freeze_panes = "A2"
    last_column = get_column_letter(len(headers))
    worksheet.auto_filter.ref = f"A1:{last_column}{max(1, worksheet.max_row)}"

    if not empty_report:
        _write_section_footer(
            worksheet,
            row_number=worksheet.max_row + 2,
            is_geography=is_geography,
            denominator=int(section.get("denominator") or 0),
            program_columns=program_columns,
        )
        if section_key == "parent_annual_income":
            _write_income_legend(worksheet, worksheet.max_row + 2)

    base_offset = 3 if is_geography else 1
    _set_capped_width(worksheet, 1, minimum=25, maximum=35)
    if is_geography:
        _set_capped_width(worksheet, 2, minimum=18, maximum=28)
        _set_capped_width(worksheet, 3, minimum=16, maximum=24)
    for index in range(len(program_columns)):
        _set_capped_width(
            worksheet,
            base_offset + index + 1,
            minimum=10,
            maximum=18,
        )
    _set_capped_width(
        worksheet,
        base_offset + len(program_columns) + 1,
        minimum=10,
        maximum=12,
    )
    _set_capped_width(
        worksheet,
        base_offset + len(program_columns) + 2,
        minimum=14,
        maximum=16,
    )
    _set_wrapped_row_heights(
        worksheet,
        first_row=1,
        last_row=worksheet.max_row,
    )
    last_column = get_column_letter(len(headers))
    worksheet.page_setup.orientation = "landscape"
    worksheet.page_setup.paperSize = worksheet.PAPERSIZE_A4
    worksheet.page_setup.scale = 100
    worksheet.page_setup.fitToWidth = 0
    worksheet.page_setup.fitToHeight = 0
    worksheet.sheet_properties.pageSetUpPr.fitToPage = False
    worksheet.page_margins.left = 0.25
    worksheet.page_margins.right = 0.25
    worksheet.page_margins.top = 0.5
    worksheet.page_margins.bottom = 0.5
    worksheet.print_title_rows = "1:1"
    worksheet.print_title_cols = "$A:$C" if is_geography else "$A:$A"
    worksheet.print_area = f"A1:{last_column}{worksheet.max_row}"


def _write_section_footer(
    worksheet: Worksheet,
    *,
    row_number: int,
    is_geography: bool,
    denominator: int,
    program_columns: list[dict[str, object]],
) -> None:
    """Write the population total below the filtered rows so every column can be verified."""

    if sum(int(program["submitted_inventory_count"]) for program in program_columns) != (
        denominator
    ):
        raise StudentProfilingWorkbookUnavailable(
            "The Student Profiling Program populations do not reconcile to the denominator."
        )
    column_number = 1
    label = _write_table_text(worksheet, row_number, column_number, "Total (profile population)")
    label.font = _BOLD
    cells = [label]
    column_number += 1
    if is_geography:
        cells.append(_write_table_text(worksheet, row_number, column_number, ""))
        cells.append(_write_table_text(worksheet, row_number, column_number + 1, ""))
        column_number += 2
    for program in program_columns:
        cells.append(
            _write_table_number(
                worksheet,
                row_number,
                column_number,
                int(program["submitted_inventory_count"]),
            )
        )
        column_number += 1
    cells.append(_write_table_number(worksheet, row_number, column_number, denominator))
    column_number += 1
    if denominator > 0:
        cells.append(
            _write_table_number(
                worksheet,
                row_number,
                column_number,
                Decimal("100.00"),
                number_format="0.00",
            )
        )
    else:
        cells.append(_write_table_text(worksheet, row_number, column_number, "N/A"))
    for cell in cells:
        cell.font = _BOLD
        cell.fill = _TOTAL_FILL
    _write_text(
        worksheet,
        row_number + 1,
        1,
        f"Denominator: n = {denominator} submitted Individual Inventories. Program columns sum "
        "to the Total row; percentages are of the profile population.",
    )


def _write_income_legend(worksheet: Worksheet, row_number: int) -> None:
    _write_text(worksheet, row_number, 1, "Parent Annual Income Brackets", bold=True)
    _write_table_header(worksheet, row_number + 1, ["Bracket", "Combined annual income"])
    row = row_number + 1
    for label, description in parent_income_bracket_legend():
        row += 1
        _write_table_text(worksheet, row, 1, label)
        _write_table_text(worksheet, row, 2, description)
    _write_text(worksheet, row + 1, 1, PARENT_INCOME_LEGEND_NOTE)


def build_student_profiling_workbook(report: dict[str, object]) -> Workbook:
    report_context = _require_dict(report.get("report_context"), "report context")
    methodology = _require_dict(report.get("methodology"), "methodology")
    coverage = _require_dict(report.get("inventory_coverage"), "Inventory Coverage")
    sections = _require_dict(report.get("sections"), "sections")
    program_columns = _program_display_columns(report.get("program_columns"))
    disclosure_warnings = _require_list(
        report.get("disclosure_warnings", []),
        "disclosure warnings",
    )
    _validate_excel_column_limit(program_columns)

    workbook = Workbook()
    _build_summary_sheet(
        workbook.active,
        report_context,
        methodology,
        coverage,
        disclosure_warnings,
        program_columns,
    )

    empty_report = int(report_context.get("submitted_inventory_count") or 0) == 0
    for section_key, sheet_name, category_header in SECTION_SPECS:
        section = _require_dict(sections.get(section_key), f"{section_key} section")
        _build_section_sheet(
            workbook,
            section_key=section_key,
            sheet_name=sheet_name,
            category_header=category_header,
            section=section,
            program_columns=program_columns,
            empty_report=empty_report,
        )

    return workbook


def _serialize_workbook(workbook: Workbook) -> bytes:
    for worksheet in workbook.worksheets:
        if worksheet.sheet_state != "visible":
            raise StudentProfilingWorkbookUnavailable(
                "The Student Profiling workbook contains an unexpected hidden worksheet."
            )
        for row in worksheet.iter_rows():
            for cell in row:
                if cell.data_type == "f":
                    raise StudentProfilingWorkbookUnavailable(
                        "The Student Profiling workbook contains an unexpected formula."
                    )
                if cell.hyperlink is not None:
                    raise StudentProfilingWorkbookUnavailable(
                        "The Student Profiling workbook contains an unexpected hyperlink."
                    )

    stream = BytesIO()
    workbook.save(stream)
    payload = stream.getvalue()
    if not payload:
        raise StudentProfilingWorkbookUnavailable(
            "The Student Profiling workbook could not be serialized."
        )
    return payload


def student_profiling_xlsx_filename(academic_year_label: str) -> str:
    safe_year = safe_report_filename_part(
        academic_year_label,
        fallback="academic-year",
    )
    return f"student-profile-{safe_year}.xlsx"


def render_student_profiling_xlsx(
    *,
    academic_year_id: UUID | None = None,
    campus_id: UUID | None = None,
    college_id: UUID | None = None,
    program_id: UUID | None = None,
    year_level: int | None = None,
    access_scope: ReportAccessScope,
) -> StudentProfilingXlsxResult:
    # Keep canonical filter/report failures outside the XLSX exception boundary.
    report = build_student_profiling_report(
        academic_year_id=academic_year_id,
        campus_id=campus_id,
        college_id=college_id,
        program_id=program_id,
        year_level=year_level,
        access_scope=access_scope,
    )

    try:
        workbook = build_student_profiling_workbook(report)
        payload = _serialize_workbook(workbook)
        report_context = _require_dict(report.get("report_context"), "report context")
        academic_year = _require_dict(
            report_context.get("academic_year"),
            "Academic Year context",
        )
        filename = student_profiling_xlsx_filename(str(academic_year.get("label") or ""))
    except StudentProfilingWorkbookUnavailable:
        raise
    except Exception as exc:
        raise StudentProfilingWorkbookUnavailable(
            "The Student Profiling report XLSX is temporarily unavailable."
        ) from exc

    return StudentProfilingXlsxResult(
        xlsx_bytes=payload,
        filename=filename,
        release_context=student_profiling_release_context(report),
    )
